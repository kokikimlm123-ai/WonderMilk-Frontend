"""Bounded, read-only adapter for James's supplied DMI/Milk workbook.

The known first two sheet layouts are checked for farm headers before any
column mapping is used. Unknown layouts fail closed. No formulas are run.
"""
import sys,json,re,math
from datetime import date,datetime
from openpyxl import load_workbook

def scalar(v):
    if isinstance(v,(datetime,date)): return v.isoformat()[:10]
    if v is None or isinstance(v,(str,int,float,bool)): return v
    return str(v)

def numeric(v):
    return isinstance(v,(float,int)) and not isinstance(v,bool) and math.isfinite(v)

def read(path,query):
    if set(query['farms'])-{'Farm 2','Farm 4','Ryokusan'}: raise ValueError('Invalid farm scope')
    cached=load_workbook(path,read_only=True,data_only=True)
    raw=load_workbook(path,read_only=True,data_only=False)
    layouts=[('feed',0,{'Farm 2':{'date':25,'cow_count':27,'delivered_kg':36,'leftover_kg':37,'dm_raw':39},'Farm 4':{'date':47,'cow_count':49,'delivered_kg':57,'leftover_kg':58,'dm_raw':60}}),
             ('milk',1,{'Farm 2':{'date':9,'cows_milked':10,'sold_kg':11,'discard_kg':12,'calf_kg':13,'source_total_kg':14,'source_average_kg':15},'Farm 4':{'date':17,'cows_milked':18,'sold_kg':19,'calf_kg':20,'discard_kg':21,'source_total_kg':22,'source_average_kg':23}})]
    output=[]
    for kind,index,mappings in layouts:
        if query.get('type') and query['type']!=kind: continue
        if len(raw.worksheets)<=index: raise ValueError('Workbook sheet layout changed')
        ws,cs=raw.worksheets[index],cached.worksheets[index]
        # Materialize a bounded area once, avoiding repeated streaming seeks.
        raw_rows=list(ws.iter_rows(min_row=1,max_row=min(ws.max_row,2200),max_col=65))
        val_rows=list(cs.iter_rows(min_row=1,max_row=min(cs.max_row,2200),max_col=65))
        for farm,columns in mappings.items():
            if farm not in query['farms']: continue
            lo=min(columns.values())-1;hi=max(columns.values())
            header=' '.join(str(c.value or '') for row in raw_rows[:20] for c in row[max(0,lo-1):hi])
            n=farm[-1]
            if not re.search(r'(?:farm|ฟาร์ม(?:ที่)?)\s*'+n+r'\b',header,re.I):
                raise ValueError('Farm header does not match the known workbook layout')
            rows=[];unknown_dates=0
            for index_row,(rr,vr) in enumerate(zip(raw_rows,val_rows),1):
                d=scalar(vr[columns['date']-1].value)
                measure=['delivered_kg','leftover_kg'] if kind=='feed' else ['sold_kg','calf_kg','discard_kg']
                # Formula-only template rows are not entered observations.
                if not any(numeric(vr[columns[k]-1].value) and rr[columns[k]-1].data_type!='f' for k in measure): continue
                if d in [None,'']: continue
                known_date=isinstance(d,str) and bool(re.fullmatch(r'\d{4}-\d{2}-\d{2}',d))
                if known_date:
                    try: date.fromisoformat(d)
                    except ValueError: known_date=False
                if known_date and d>query['today']: continue
                if not known_date:
                    unknown_dates+=1
                    if query.get('date_from') or query.get('date_to'): continue
                if known_date and query.get('date_from') and d<query['date_from']: continue
                if known_date and query.get('date_to') and d>query['date_to']: continue
                record={k:scalar(vr[col-1].value) for k,col in columns.items()}
                record.update(source_row=index_row,issues=[])
                if not known_date: record['issues'].append('Unresolved date format; retain original label')
                if kind=='milk':
                    if all(numeric(record[k]) for k in measure):
                        record['calculated_complete_total_kg']=sum(record[k] for k in measure)
                        if numeric(record['source_total_kg']) and abs(record['source_total_kg']-record['calculated_complete_total_kg'])>0.1:
                            record['issues'].append('Source total does not match components')
                    else: record['issues'].append('Missing component(s); cannot calculate complete farm milk total')
                else:
                    dm=vr[columns['dm_raw']-1]
                    if numeric(dm.value) and '%' in dm.number_format: record['dm_percent']=dm.value*100
                    elif dm.value is not None: record['issues'].append('DM raw value retained; percentage/fraction unit requires confirmation')
                rows.append(record)
            output.append({'farm':farm,'type':kind,'tab':ws.title,'matched_rows':len(rows),'presented_rows':rows[-40:],'unresolved_date_rows':unknown_dates,'limit_note':'At most 40 last populated source rows; row order alone is not proof of chronology.'})
    cached.close();raw.close()
    return {'tables':output}

if __name__=='__main__':
    try: print(json.dumps(read(sys.argv[1],json.loads(sys.argv[2])),ensure_ascii=False,allow_nan=False))
    except Exception as exc:
        print(json.dumps({'error':type(exc).__name__}),file=sys.stderr)
        sys.exit(1)
