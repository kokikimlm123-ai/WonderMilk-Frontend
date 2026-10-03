"""Read-only Ryokusan source adapter. Preserve cached values and original labels."""
import sys,json,re
from datetime import datetime,date
from openpyxl import load_workbook
def scalar(v):
    if isinstance(v,(date,datetime)):return v.isoformat()[:10]
    return v if v is None or isinstance(v,(str,int,float,bool)) else str(v)
def day(v):
    if isinstance(v,(date,datetime)):return v.isoformat()[:10]
    if isinstance(v,str):
        for fmt in ('%d/%m/%Y','%d/%m/%y','%d/%m/%y %H:%M','%Y-%m-%d'):
            try:return datetime.strptime(v,fmt).date().isoformat()
            except ValueError:pass
    return None
def read(path,q):
    if q['farms']!=['Ryokusan']:raise ValueError('Invalid farm scope')
    w=load_workbook(path,read_only=True,data_only=True)
    if w['Visit_Record']['C2'].value!='Ryokusan Farm':raise ValueError('Wrong farm workbook')
    specs=[
      ('milk','Average milk and milk sold reco',2,8,0,None),
      ('feed','Feeding record',4,17,0,None),
      ('cow','Cow_Data',1,15,None,0),
      ('health','Injection and treatment record',1,7,1,0),
      ('health','CMT Test result',4,9,1,2),
      ('cow','Insemination Record',1,16,3,0),
      ('cow','Pregnancy check Record',1,9,2,0),
      ('milk_test','Milk Quality Test Result',5,11,None,None),
      ('general','Visit_Record',5,4,None,None),
      ('general','Inspection Summary',1,12,None,None),
      ('general','Medicine Inventory record',3,4,None,None),
      ('general','Synchronization Program',1,12,None,None)]
    out=[]
    for kind,name,header,ncols,datecol,idcol in specs:
        if q.get('type') and kind!=q['type'] and not(kind=='general' and q['type']=='health'):continue
        s=w[name]
        rows=list(s.iter_rows(max_row=min(s.max_row or 2200,2200),max_col=ncols,values_only=True))
        selected=[]
        for n,row in enumerate(rows[header:],header+1):
            if not any(v is not None for v in row):continue
            if q.get('cow_id') and idcol is not None and str(row[idcol])!=str(q['cow_id']):continue
            d=day(row[datecol]) if datecol is not None else None
            if datecol is not None:
                if not d:continue
                if q.get('date_from') and d<q['date_from']:continue
                if q.get('date_to') and d>q['date_to']:continue
            selected.append({'row':n,'date':d,'cells':[scalar(v) for v in row]})
        if datecol is not None:selected.sort(key=lambda r:r['date'])
        limit=(100 if kind in ('cow','general') else 40) if q.get('type') else 15
        out.append({'farm':'Ryokusan','type':kind,'tab':name,'headers':[[scalar(v) for v in row] for row in rows[:header]],'matched_rows':len(selected),'records':selected[-limit:],'truncated':len(selected)>limit})
    w.close()
    return {'farm':'Ryokusan','tables':out,'notes':[
      'Source snapshot supplied 2026-10-03. Sheet dates may differ from the inspection date. Do not recompute synchronization schedules.',
      'Values are cached source values. Missing values and formula errors are not zero. Preserve units as labeled; do not infer SCC multiplier.',
      'Milk sheet contains herd AM/PM sold milk, calf milk, source total and average. There is no separate discarded-milk column.',
      'Milking summary is not mapped: its headers do not align with the supplied data. Do not infer individual milk yields from it.',
      'Cow dates and DIM may conflict; report source values without silently correcting them.']}
if __name__=='__main__':
    print(json.dumps(read(sys.argv[1],json.loads(sys.argv[2])),ensure_ascii=False,allow_nan=False))
